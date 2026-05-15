#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from docx import Document
import json
import re

doc = Document(r'd:\codeX_project\English_logic_map\ELP\初中语法\句子的种类（教师版）.docx')

all_content = []

current_topic = None
current_subtopic = None

for para in doc.paragraphs:
    text = para.text.strip()
    if not text:
        continue

    if text.startswith('★★★考点'):
        topic = text.replace('★★★考点', '').replace('★', '').strip()
        current_topic = topic
        current_subtopic = None
        all_content.append({'type': 'topic', 'topic': topic, 'content': []})
        continue

    if text.startswith('【一、') or text.startswith('【二、') or text.startswith('【三、') or text.startswith('【四、'):
        if all_content and all_content[-1]['type'] == 'topic':
            all_content[-1]['subtopic'] = text
        current_subtopic = text
        continue

    if re.match(r'^\d+\.', text) and len(text) < 300:
        if all_content:
            all_content[-1]['content'].append({'type': 'example', 'text': text})
        continue

    if text.startswith('如：') or text.startswith('比如'):
        if all_content:
            all_content[-1]['content'].append({'type': 'note', 'text': text})
        continue

    if len(text) > 5:
        if all_content:
            all_content[-1]['content'].append({'type': 'text', 'text': text})

tables_data = []
for i, table in enumerate(doc.tables):
    table_data = []
    for row in table.rows:
        row_text = [cell.text.strip() for cell in row.cells]
        table_data.append(row_text)
    tables_data.append({'table_index': i, 'data': table_data})

result = {
    'document': '句子的种类（教师版）',
    'paragraphs': all_content,
    'tables': tables_data
}

output_path = r'd:\codeX_project\English_logic_map\ELP\output\sentence_types_full.json'
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)

print(f"已保存 {len(all_content)} 个主题")
print(f"已保存 {len(tables_data)} 个表格")
print(f"输出文件: {output_path}")