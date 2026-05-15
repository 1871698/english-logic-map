#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from docx import Document
import json
import re

doc = Document(r'd:\codeX_project\English_logic_map\ELP\初中语法\句子的种类（教师版）.docx')

sentence_types = {
    "陈述句": [],
    "疑问句": {
        "一般疑问句": [],
        "特殊疑问句": [],
        "选择疑问句": [],
        "反意疑问句": []
    },
    "感叹句": [],
    "祈使句": []
}

current_type = None
current_subtype = None
current_content = []

for para in doc.paragraphs:
    text = para.text.strip()
    if not text:
        continue

    if text.startswith('★★★考点'):
        if '反意疑问句' in text:
            current_type = "疑问句"
            current_subtype = "反意疑问句"
        elif '感叹句' in text:
            current_type = "感叹句"
            current_subtype = None
        elif '祈使句' in text:
            current_type = "祈使句"
            current_subtype = None
        elif '一般疑问句' in text:
            current_type = "疑问句"
            current_subtype = "一般疑问句"
        elif '特殊疑问句' in text:
            current_type = "疑问句"
            current_subtype = "特殊疑问句"
        elif '选择疑问句' in text:
            current_type = "疑问句"
            current_subtype = "选择疑问句"
        elif '陈述句' in text:
            current_type = "陈述句"
            current_subtype = None
        continue

    if text.startswith('【一、') and '定义' in text:
        current_content.append({'type': 'definition', 'text': text})
        continue

    if text.startswith('【二、') or text.startswith('【三、') or text.startswith('【四、'):
        if current_content:
            save_content(current_type, current_subtype, current_content)
            current_content = []
        current_content.append({'type': 'rule', 'text': text})
        continue

    if re.match(r'^\d+\.', text) and len(text) < 200:
        current_content.append({'type': 'example', 'text': text})
        continue

    if text.startswith('如：') or text.startswith('比如'):
        current_content.append({'type': 'note', 'text': text})
        continue

    if '考点' in text or '注意' in text or '说明' in text:
        current_content.append({'type': 'note', 'text': text})
        continue

    if current_content and len(text) > 10:
        current_content.append({'type': 'explanation', 'text': text})

def save_content(type_name, subtype, content):
    if not content:
        return
    item = {'rules': [], 'examples': [], 'notes': []}
    for c in content:
        if c['type'] == 'rule':
            item['rules'].append(c['text'])
        elif c['type'] == 'example':
            item['examples'].append(c['text'])
        elif c['type'] == 'note':
            item['notes'].append(c['text'])
        elif c['type'] == 'explanation':
            item['notes'].append(c['text'])

    if subtype:
        if subtype not in sentence_types[type_name]:
            sentence_types[type_name][subtype] = []
        sentence_types[type_name][subtype].append(item)
    else:
        if isinstance(sentence_types[type_name], list):
            sentence_types[type_name].append(item)
        else:
            sentence_types[type_name] = item

save_content(current_type, current_subtype, current_content)

for i, table in enumerate(doc.tables):
    table_data = []
    for row in table.rows:
        row_text = [cell.text.strip() for cell in row.cells]
        table_data.append(row_text)
    print(f"\n表格 {i+1}:")
    for row in table_data[:5]:
        print(row)

print("\n\n=== 结构化数据 ===")
print(json.dumps(sentence_types, ensure_ascii=False, indent=2))