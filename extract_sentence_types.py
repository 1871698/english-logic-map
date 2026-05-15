#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from docx import Document
import json

doc = Document(r'd:\codeX_project\English_logic_map\ELP\初中语法\句子的种类（教师版）.docx')

print("=== 文档内容提取 ===\n")

current_section = None
sections = {}

for para in doc.paragraphs:
    text = para.text.strip()
    if not text:
        continue

    # 检测是否是标题（通常以特定格式开头）
    if text.startswith(('一、', '二、', '三、', '四、', '五、', '六、', '七、', '八、', '九、', '十、')):
        current_section = text
        sections[current_section] = []
        print(f"\n【{text}】")
    elif text.startswith('（') and '）' in text and len(text) < 50:
        # 子标题
        if current_section:
            sections[current_section].append({'type': 'subsection', 'text': text})
        print(f"  - {text}")
    elif text.startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.', '10.')):
        # 列表项
        if current_section:
            sections[current_section].append({'type': 'example', 'text': text})
        print(f"    {text}")
    elif text and current_section:
        sections[current_section].append({'type': 'content', 'text': text})
        print(f"  {text}")

# 表格提取
print("\n\n=== 表格内容 ===")
for i, table in enumerate(doc.tables):
    print(f"\n表格 {i+1}:")
    for row in table.rows:
        row_text = [cell.text.strip() for cell in row.cells]
        print(" | ".join(row_text))

# 保存为结构化JSON
output = {
    "document_name": "句子的种类（教师版）",
    "sections": sections
}

with open(r'd:\codeX_project\English_logic_map\ELP\output\sentence_types_raw.json', 'w', encoding='utf-8') as f:
    json.dump(output, f, ensure_ascii=False, indent=2)

print("\n\n已保存到 output/sentence_types_raw.json")